#!/usr/bin/env python3
"""Build the FIT3143 Applied #2 deck (v3): minimal blue, bullets only, 30pt+ body.

Run:  python3 build_deck.py   ->  FIT3143_Applied2_Slides_v3.pptx
Every body text box is auto-checked: if text at the chosen size would overflow
its box, the script prints a WARNING (so nothing silently runs off a slide).
"""
import os, re
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
T1 = os.path.join(ROOT, "task1")
T2 = os.path.join(ROOT, "task2")

NAVY = RGBColor(0x0B, 0x2A, 0x5B)
BLUE = RGBColor(0x25, 0x63, 0xEB)
SKY = RGBColor(0xEA, 0xF1, 0xFF)
MID = RGBColor(0xBF, 0xD4, 0xFF)
INK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
_fc = {}
WARN = []

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- text fitting ----------
def _font(size, bold):
    k = (size, bold)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(BOLD if bold else REG, int(size * 10))
    return _fc[k]


def n_lines(text, size, width_in, bold=False):
    f = _font(size, bold)
    maxw = width_in * 72 * 10  # points*10
    lines = 0
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        cnt = 1
        for w in words:
            t = (cur + " " + w).strip()
            if f.getlength(t) <= maxw:
                cur = t
            else:
                cnt += 1
                cur = w
        lines += cnt
    return lines


def check(label, items, size, w_in, h_in, bold=False, bullet=True, gap=8):
    ind = 0.4 if bullet else 0
    total = 0
    for it in items:
        total += n_lines(it, size, w_in - ind - 0.2, bold) * size * 1.2 + gap
    total -= gap
    if total / 72 > h_in * 1.06:
        WARN.append(f"{label}: needs {total/72:.2f}in, box {h_in:.2f}in")


# ---------- primitives ----------
def rect(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.08, dash=False, lw=2):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            ln = s.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")
    s.shadow.inherit = False
    return s


def _run(p, text, size, bold=False, color=INK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return r


def _bullet(p, color=BLUE):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.4))))
    pPr.set("indent", str(-int(Inches(0.4))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    c = etree.SubElement(pPr, qn("a:buClr"))
    sc = etree.SubElement(c, qn("a:srgbClr"))
    sc.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar"))
    bc.set("char", "•")


def text(sl, x, y, w, h, content, size=30, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, label="",
         gap=8, bcolor=BLUE):
    items = content if isinstance(content, list) else [content]
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = anchor
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        _run(p, it, size, bold, color)
        if bullets:
            _bullet(p, bcolor)
    check(label or items[0][:30], items, size, w, h, bold, bullets, gap)
    return tb


def label_in(shape, content, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
             label=""):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = content if isinstance(content, list) else [content]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _run(p, ln, size, bold, color)
    w = shape.width / 914400
    h = shape.height / 914400
    check(label or lines[0][:30], lines, size, w, h, bold, False, 2)


def picture(sl, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    rect(sl, x, y, w, h, fill=WHITE, line=MID, lw=1.5)
    sl.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


def new_slide(title, presenter, sources="", notes="", appendix=False):
    sl = prs.slides.add_slide(BLANK)
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    rect(sl, 0, 0, 0.28, 7.5, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 0.6, 0.3, 11.2, 1.0, title, size=40, bold=True, color=NAVY,
         anchor=MSO_ANCHOR.MIDDLE, label="TITLE " + title, gap=0)
    if appendix:
        pill = rect(sl, 0.6, 1.2, 12.1, 0.42, fill=SKY, radius=0.5)
        label_in(pill, "APPENDIX: extra material for Q&A, not part of the 7-minute talk",
                 size=16, bold=False, color=NAVY, label="appx pill")
    else:
        tag = rect(sl, 10.6, 0.15, 2.5, 0.38, fill=SKY, radius=0.5)
        label_in(tag, "Presenter: " + presenter, size=14, bold=True, color=NAVY,
                 label="presenter tag")
    if sources:
        text(sl, 0.6, 7.0, 12.2, 0.4, "Sources: " + sources, size=12, color=GREY,
             label="sources", gap=0)
    sl.notes_slide.notes_text_frame.text = notes
    return sl


def code_frame(sl, x, y, w, h, n, what, lines):
    rect(sl, x, y, w, h, fill=SKY, line=BLUE, dash=True, lw=2.5)
    text(sl, x + 0.2, y + 0.15, w - 0.4, h - 0.3,
         [f"YOUR SCREENSHOT {n}", what, lines], size=20, bold=False, color=NAVY,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="code frame",
         gap=6)


def stat(sl, x, y, w, h, big, small, fill=SKY, bigsize=54):
    rect(sl, x, y, w, h, fill=fill)
    text(sl, x, y + 0.05, w, h * 0.55, big, size=bigsize, bold=True, color=BLUE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="stat " + big, gap=0)
    text(sl, x, y + h * 0.55, w, h * 0.42, small, size=30, color=INK,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, label="stat label " + small,
         gap=0)


def box(sl, x, y, w, h, content, size=30, fill=BLUE, color=WHITE):
    s = rect(sl, x, y, w, h, fill=fill)
    label_in(s, content, size=size, color=color, label="box " + str(content)[:20])
    return s


def arrow(sl, x, y, w, h):
    a = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = MID
    a.line.fill.background()
    return a


# =====================================================================
# 1  Title
# =====================================================================
sl = prs.slides.add_slide(BLANK)
sl.background.fill.solid()
sl.background.fill.fore_color.rgb = NAVY
rect(sl, 0, 0, 0.4, 7.5, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
text(sl, 0.9, 0.6, 11.5, 0.5, "Monash FIT3143 Parallel Computing  |  Applied #2",
     size=24, color=MID, label="t1 unit")
text(sl, 0.9, 1.3, 11.6, 2.4,
     "GPU Image Rotation and the Ethics of Scaling HPC for AI", size=54, bold=True,
     color=WHITE, label="t1 title", gap=0)
text(sl, 0.9, 3.9, 11.6, 1.4,
     ["Erwyna Soo Wen Xin (36555789)", "Taabish Farooq Bhat (35473932)"], size=32,
     bold=True, color=WHITE, label="t1 names", gap=4)
text(sl, 0.9, 5.2, 11.6, 0.5,
     "esoo0013@student.monash.edu   |   ttaa0006@student.monash.edu", size=22,
     color=MID, label="t1 emails")
for i, t in enumerate(["Task 1: GPU design", "Task 2: Ethics of HPC",
                       "Limits and Q&A"]):
    s = rect(sl, 0.9 + i * 4.0, 6.0, 3.8, 0.8, fill=BLUE, radius=0.5)
    label_in(s, t, size=24, color=WHITE, label="agenda chip")
sl.notes_slide.notes_text_frame.text = (
    "ERWYNA (0:12)\n"
    "Hi, I'm Erwyna and this is Taabish.\n"
    "Taabish covers Task 1: rotating big images on a GPU with CUDA.\n"
    "I cover Task 2: the ethics of scaling HPC for AI, then our limits.")

# =====================================================================
# 2  Task 1 intro
# =====================================================================
sl = new_slide("Task 1: 33 million independent pixels", "Taabish",
               sources="8K = 7680 x 4320.",
               notes="TAABISH (0:20)\n"
               "An 8K image has about 33 million pixels.\n"
               "Each output pixel depends only on the input, so this is data parallel.\n"
               "We use inverse mapping: one thread per output pixel.\n"
               "Each thread rotates back to find its source pixel. One writer per pixel, no holes.")
text(sl, 0.6, 1.5, 6.0, 5.2,
     ["8K image: 33,177,600 pixels", "Every pixel is independent",
      "So the job is data parallel", "One thread per output pixel",
      "One writer per pixel, no holes"], size=32, bullets=True, label="s2 bullets",
     gap=14)
box(sl, 7.1, 1.8, 5.6, 1.1, "Output pixel (x, y)", size=30)
box(sl, 7.1, 3.35, 5.6, 1.1, "Rotate back (R transpose)", size=30, fill=NAVY)
box(sl, 7.1, 4.9, 5.6, 1.1, "Source pixel", size=30)
for yy in (2.95, 4.5):
    a = sl.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.65), Inches(yy), Inches(0.5), Inches(0.4))
    a.fill.solid(); a.fill.fore_color.rgb = MID; a.line.fill.background()
text(sl, 7.1, 6.1, 5.6, 0.6, "Inverse mapping", size=30, bold=True, color=NAVY,
     align=PP_ALIGN.CENTER, label="s2 caption")

# =====================================================================
# 3  1a data path
# =====================================================================
sl = new_slide("1a. Host to device and back", "Taabish",
               sources="NVIDIA (n.d.-d, 2026a, 2026b). Measured on Tesla T4.",
               notes="TAABISH (0:35)\n"
               "cudaMalloc reserves both buffers in GPU memory.\n"
               "For host to device, the GPU's DMA copy engine pulls the bytes over PCIe.\n"
               "From pageable memory, the driver first copies into a pinned staging buffer.\n"
               "cudaMallocHost gives pinned memory, so the DMA reads it directly.\n"
               "The kernel works only in GPU memory, then device to host brings the result back.\n"
               "PCIe is the narrow pipe: pinned ran at 12.3 GB per second, GPU memory is 320.")
box(sl, 0.6, 1.5, 3.4, 1.5, ["CPU memory", "(host)"], size=30)
box(sl, 9.3, 1.5, 3.4, 1.5, ["GPU memory", "(device)"], size=30)
arrow(sl, 4.1, 1.85, 5.1, 0.8)
text(sl, 4.1, 2.7, 5.1, 0.5, "PCIe: 12.3 GB/s", size=30, bold=True, color=BLUE,
     align=PP_ALIGN.CENTER, label="pcie lbl")
text(sl, 8.8, 3.0, 4.4, 0.6, "GDDR6: 320 GB/s", size=30, bold=True, color=NAVY,
     align=PP_ALIGN.CENTER, label="gddr lbl")
text(sl, 0.6, 3.9, 12.1, 3.0,
     ["Pageable memory: extra staging copy, 4.6 GB/s",
      "Pinned memory: direct DMA, 12.3 GB/s",
      "The kernel only works inside GPU memory",
      "Unified Memory and zero-copy move data on demand"],
     size=30, bullets=True, label="s3 bullets", gap=10)

# =====================================================================
# 4  CUDA model
# =====================================================================
sl = new_slide("1b. CUDA model", "Taabish",
               sources="NVIDIA (2026b).",
               notes="TAABISH (0:25)\n"
               "The host allocates, copies, picks the launch size, launches the kernel, then syncs and checks errors.\n"
               "The device runs thousands of threads, SIMT style.\n"
               "Our block is 16 by 16: 256 threads, or 8 warps of 32.\n"
               "For 8K the grid is 480 by 270 blocks, one thread per pixel.")
text(sl, 0.6, 1.5, 7.2, 5.3,
     ["Host (CPU): allocate, copy, launch, sync",
      "Device (GPU): thousands of SIMT threads",
      "Block: 16 x 16 = 256 threads = 8 warps",
      "Grid: 480 x 270 = 129,600 blocks"],
     size=32, bullets=True, label="s4 bullets", gap=16)
rect(sl, 8.2, 1.6, 4.6, 3.7, fill=SKY, line=BLUE, lw=2)
for r_ in range(4):
    for c_ in range(5):
        hl = (r_ == 1 and c_ == 2)
        rect(sl, 8.4 + c_ * 0.84, 1.8 + r_ * 0.84, 0.74, 0.74,
             fill=BLUE if hl else WHITE, line=BLUE, radius=0.1, lw=1.5)
text(sl, 8.2, 5.35, 4.6, 0.55, "Grid of blocks", size=30, bold=True, color=NAVY,
     align=PP_ALIGN.CENTER, label="s4 cap1")
text(sl, 8.2, 5.95, 4.6, 0.55, "Block: 16 x 16 threads", size=30, color=INK,
     align=PP_ALIGN.CENTER, label="s4 cap2")

# =====================================================================
# 5  Code proof
# =====================================================================
sl = new_slide("1b. Our kernel and launch", "Taabish",
               sources="rotate_cuda.cu in our submission zip.",
               notes="TAABISH (0:20)\n"
               "This is the kernel from our submission.\n"
               "Each thread finds its pixel from blockIdx times blockDim plus threadIdx.\n"
               "The guard drops threads past the edge.\n"
               "It rotates back to the source pixel, so every output pixel has one writer.\n"
               "The launch rounds the grid up, and we check the error straight after, because launches are asynchronous.")
code_frame(sl, 0.6, 1.4, 6.0, 3.4, 1, "Kernel rotate_nn_2d", "rotate_cuda.cu lines 145 to 158")
code_frame(sl, 6.8, 1.4, 6.0, 3.4, 2, "Grid size and launch", "rotate_cuda.cu lines 360 to 383")
text(sl, 0.6, 5.0, 12.2, 1.9,
     ["Thread finds its pixel from blockIdx, blockDim, threadIdx",
      "Guard drops extra threads; grid is rounded up",
      "Launch is async, so we check the error right after"],
     size=30, bullets=True, label="s5 bullets", gap=6)

# =====================================================================
# 6  Results
# =====================================================================
sl = new_slide("1b. Each feature, measured", "Taabish",
               sources="Tesla T4, 8K at 30 degrees, v2 nearest, mean of 10 runs. CPU baseline: one core.",
               notes="TAABISH (0:40)\n"
               "We built one version per CUDA feature: 1D grid, 2D blocks, bilinear, pageable, streams, texture.\n"
               "The 2D grid beat the 1D grid, 1.55 against 2.96 milliseconds.\n"
               "Against one CPU core the kernel is 328 times faster.\n"
               "But end to end it is only 29.6 times, because copies take 91 percent of the time.\n"
               "Amdahl caps us at 32.6 times.\n"
               "A 14-thread laptop beats the T4 end to end. So the kernel is memory bound.")
picture(sl, os.path.join(T1, "results", "g1_speedup.png"), 0.6, 1.5, 7.4, 3.3)
text(sl, 0.6, 4.95, 7.4, 2.0,
     ["6 versions built, one per feature", "2D grid beat 1D: 1.55 vs 2.96 ms",
      "Amdahl cap: 32.6x"], size=30, bullets=True, label="s6 bullets", gap=6)
stat(sl, 8.3, 1.5, 4.5, 1.6, "328x", "kernel only")
stat(sl, 8.3, 3.25, 4.5, 1.6, "29.6x", "end to end")
stat(sl, 8.3, 5.0, 4.5, 1.6, "91%", "time is copies", bigsize=48)

# =====================================================================
# 7  GDS
# =====================================================================
sl = new_slide("1c. GPUDirect Storage", "Taabish",
               sources="Thompson & Newburn (2019); NVIDIA (n.d.-a).",
               notes="TAABISH (0:30)\n"
               "GPUDirect Storage lets NVMe DMA straight into GPU memory.\n"
               "The CPU only keeps the control path, so the host bounce buffer goes away.\n"
               "For the company: yes, but only for the bulk pipeline, and only if frames are raw or decoded on the GPU with nvJPEG.\n"
               "It does not help one photo, JPEGs decoded on the CPU, or images already in RAM.")
rect(sl, 0.6, 1.5, 6.0, 4.4, fill=SKY)
text(sl, 0.8, 1.6, 5.6, 0.7, "Helps", size=34, bold=True, color=BLUE, label="s7 h1")
text(sl, 0.8, 2.4, 5.6, 3.4,
     ["Bulk raw frames", "JPEGs decoded on the GPU (nvJPEG)", "Disk I/O is the bottleneck"],
     size=30, bullets=True, label="s7 b1", gap=12)
rect(sl, 6.8, 1.5, 6.0, 4.4, fill=SKY)
text(sl, 7.0, 1.6, 5.6, 0.7, "Does not help", size=34, bold=True, color=NAVY, label="s7 h2")
text(sl, 7.0, 2.4, 5.6, 3.4,
     ["One photo", "JPEGs decoded on the CPU", "Images already in RAM"],
     size=30, bullets=True, label="s7 b2", gap=12, bcolor=NAVY)
text(sl, 0.6, 6.1, 12.2, 0.7, "Not tested: Colab has no nvidia-fs module", size=30,
     bold=True, color=NAVY, label="s7 note")

# =====================================================================
# 8  Energy
# =====================================================================
sl = new_slide("2a. Data centre demand keeps climbing", "Erwyna",
               sources="International Energy Agency (IEA, 2025).",
               notes="ERWYNA (0:15)\n"
               "The IEA says data centres used 415 terawatt hours in 2024.\n"
               "That could reach about 1,200 by 2035, and GPU servers drive the growth.")
picture(sl, os.path.join(T2, "charts", "C2_iea_datacentre_electricity.png"), 0.6, 1.5, 7.4, 4.5)
stat(sl, 8.3, 1.5, 4.5, 1.6, "415 TWh", "in 2024", bigsize=44)
stat(sl, 8.3, 3.25, 4.5, 1.6, "945 TWh", "by 2030", bigsize=44)
stat(sl, 8.3, 5.0, 4.5, 1.6, "1,200 TWh", "by 2035 (about)", bigsize=44)
text(sl, 0.6, 6.15, 7.4, 0.7, "GPU servers drive the growth", size=30, bold=True,
     color=NAVY, label="s8 caption")

# =====================================================================
# 9  Same chip different grid
# =====================================================================
sl = new_slide("2a. Same chip, different grid", "Erwyna",
               sources="TOP500 (2026); LUMI consortium (2023).",
               notes="ERWYNA (0:20)\n"
               "Frontier and LUMI use the same MI250X nodes, and their GFLOPS per watt are almost equal.\n"
               "But LUMI runs on hydropower and heats part of a Finnish town.\n"
               "So a FLOP's carbon cost depends on the grid, not only the chip.")
rect(sl, 0.6, 1.5, 6.0, 4.1, fill=SKY)
text(sl, 0.8, 1.6, 5.6, 0.7, "Frontier (USA)", size=34, bold=True, color=NAVY, label="s9 h1")
text(sl, 0.8, 2.4, 5.6, 1.3, "55.0", size=72, bold=True, color=BLUE, label="s9 n1", gap=0)
text(sl, 0.8, 3.7, 5.6, 0.6, "GFLOPS per watt", size=30, label="s9 u1")
rect(sl, 6.8, 1.5, 6.0, 4.1, fill=SKY)
text(sl, 7.0, 1.6, 5.6, 0.7, "LUMI (Finland)", size=34, bold=True, color=NAVY, label="s9 h2")
text(sl, 7.0, 2.4, 5.6, 1.3, "53.4", size=72, bold=True, color=BLUE, label="s9 n2", gap=0)
text(sl, 7.0, 3.7, 5.6, 0.6, "GFLOPS per watt", size=30, label="s9 u2")
text(sl, 7.0, 4.3, 5.6, 1.2, ["100% hydropower", "Waste heat warms a town"], size=30,
     bullets=True, label="s9 lumi", gap=2)
bar = rect(sl, 0.6, 5.85, 12.2, 1.0, fill=NAVY, radius=0.2)
label_in(bar, "Carbon per FLOP depends on the grid, not only the chip", size=30,
         label="s9 bar")

# =====================================================================
# 10  Malaysia + myth
# =====================================================================
sl = new_slide("2a. Malaysia lens and a myth check", "Erwyna",
               sources="Ember (2026); Loo (2025); Bernama (2026); LUMI consortium (2023); Jamison et al. (2025); Giles (2025); IEA (2025).",
               notes="ERWYNA (0:22)\n"
               "Johor went from 10 megawatts to over a gigawatt in about three years.\n"
               "Data centres could use almost a third of Peninsular electricity by 2035, on a grid that is 79 percent fossil.\n"
               "Kajaani in Finland is the opposite: hydropower, cold air, and waste heat for the town.\n"
               "One myth check: AI does not emit more CO2 than aviation today. That is a 2030 projection.\n"
               "(If asked: Johor also stopped approving new water-hungry Tier 1 and 2 sites.)")
rect(sl, 0.6, 1.5, 6.0, 3.7, fill=SKY)
text(sl, 0.8, 1.55, 5.6, 0.7, "Johor, Malaysia", size=34, bold=True, color=NAVY, label="s10 h1")
text(sl, 0.8, 2.3, 5.6, 2.9,
     ["79% fossil grid", "10 MW to 1.3 GW in 3 years", "Up to 31% of Peninsular power by 2035"],
     size=30, bullets=True, label="s10 b1", gap=8)
rect(sl, 6.8, 1.5, 6.0, 3.7, fill=SKY)
text(sl, 7.0, 1.55, 5.6, 0.7, "Kajaani, Finland", size=34, bold=True, color=NAVY, label="s10 h2")
text(sl, 7.0, 2.3, 5.6, 2.9,
     ["100% hydropower", "Cold climate cooling", "Waste heat warms the town"],
     size=30, bullets=True, label="s10 b2", gap=8)
bar = rect(sl, 0.6, 5.4, 12.2, 1.45, fill=NAVY, radius=0.15)
text(sl, 0.8, 5.45, 11.8, 1.35,
     ["Myth: \"AI out-emits aviation\"", "It is a 2030 projection. Today: about 180 Mt CO2"],
     size=30, bold=False, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, label="s10 myth", gap=2)

# =====================================================================
# 11  Governance
# =====================================================================
sl = new_slide("2b. Governance, security and fairness", "Erwyna",
               sources="Pilz et al. (2025); Zhu et al. (2019); BIS (2022); MITI (2025); Common Crawl (2026); Ahuja et al. (2023).",
               notes="ERWYNA (0:35)\n"
               "Where the GPUs are decides which country's law covers the data.\n"
               "In data-parallel training, nodes swap gradients, and gradients can leak training data. So data residency becomes a scheduling rule.\n"
               "Security is dual-use: the same chips train a hospital model or a military model. The US limits exports, and since July 2025 Malaysia needs a permit.\n"
               "Fairness has a compute cost too: Malay is under 0.1 percent of web text.")
rows = [("Governance", ["75% of AI supercomputer power is in the US",
                        "Gradients can leak data: keep it local"]),
        ("Security", ["Dual-use chips: hospital or military",
                      "US export limits; Malaysia needs permits"]),
        ("Fairness", ["Malay is 0.086% of web text",
                      "Tamil needs about 10 tokens per word"])]
for i, (h_, bl) in enumerate(rows):
    y = 1.5 + i * 1.8
    rect(sl, 0.6, y, 12.2, 1.65, fill=SKY)
    c = rect(sl, 0.6, y, 2.9, 1.65, fill=BLUE)
    label_in(c, h_, size=30, label="s11 chip " + h_)
    text(sl, 3.6, y + 0.05, 9.1, 1.55, bl, size=30, bullets=True,
         anchor=MSO_ANCHOR.MIDDLE, label="s11 " + h_, gap=4)

# =====================================================================
# 12  Compute divide
# =====================================================================
sl = new_slide("2c. The compute divide", "Erwyna",
               sources="Stanford HAI (2026); Epoch AI (2026); Khandelwal et al. (2025); NCI (n.d.-a, n.d.-b).",
               notes="ERWYNA (0:25)\n"
               "Over 90 percent of notable 2025 models came from industry.\n"
               "Frontier training compute grows about five times a year, and the best academic model sits about three thousand times lower.\n"
               "Most academics have one to eight GPUs. (If asked: NCI Gadi has 776 GPUs and demand near three times its allocation.)\n"
               "If only companies can train frontier models, only companies can audit them.")
picture(sl, os.path.join(T2, "charts", "C3_epoch_training_compute.png"), 0.6, 1.5, 7.4, 4.5)
stat(sl, 8.3, 1.5, 4.5, 1.6, ">90%", "of models: industry", bigsize=48)
stat(sl, 8.3, 3.25, 4.5, 1.6, "5x / year", "frontier compute", bigsize=44)
stat(sl, 8.3, 5.0, 4.5, 1.6, "1 to 8", "academic GPUs", bigsize=48)
text(sl, 0.6, 6.15, 7.4, 0.7, "Only companies can audit", size=30,
     bold=True, color=NAVY, label="s12 caption")

# =====================================================================
# 13  SCALE
# =====================================================================
sl = new_slide("2d. Our SCALE checklist", "Erwyna",
               sources="UNESCO (2021); OECD (2024); MOSTI (2024); European Parliament & Council (2024); Wiesner et al. (2021).",
               notes="ERWYNA (0:35)\n"
               "UNESCO, the OECD and Malaysia's AIGE give values, but no number to report.\n"
               "The EU is closer: the AI Act flags models above 10 to the 25 FLOP, and big sites must report PUE.\n"
               "So we propose SCALE. Each step is a Slurm setting and a reported number.\n"
               "Size right. Carbon-aware: jobs that can wait cut CO2 by up to 19 percent. Account every joule. Lawful, local data. Equitable access.")
scale = [("S", "Size right: stop when efficiency drops"),
         ("C", "Carbon-aware: run flexible jobs at low-carbon times"),
         ("A", "Account every joule: report energy per job"),
         ("L", "Lawful, local data: keep data in its country"),
         ("E", "Equitable access: fair-share queues")]
for i, (l_, t_) in enumerate(scale):
    y = 1.45 + i * 0.98
    rect(sl, 0.6, y, 12.2, 0.88, fill=SKY)
    c = rect(sl, 0.6, y, 0.95, 0.88, fill=BLUE)
    label_in(c, l_, size=36, label="scale letter")
    text(sl, 1.7, y, 11.0, 0.88, t_, size=30, anchor=MSO_ANCHOR.MIDDLE,
         label="scale " + l_, gap=0)
bar = rect(sl, 0.6, 6.4, 12.2, 0.55, fill=NAVY, radius=0.3)
label_in(bar, "Principles give values. The EU gives numbers.", size=30, bold=False,
         label="s13 bar")

# =====================================================================
# 14  Limits
# =====================================================================
sl = new_slide("Limitations, future work and takeaways", "Erwyna",
               sources="Epoch AI (2026); Radovanovic et al. (2023).",
               notes="ERWYNA (0:28)\n"
               "Task 1 limits: a one-core CPU baseline on a shared Colab VM, so our speed-ups belong to that setup, and we could not test GDS.\n"
               "Next: nvJPEG, GDS and streams on a GDS server.\n"
               "Task 2 limits: most model energy figures are estimates, so we measured ours: about 0.7 joules per 8K image on the T4.\n"
               "Google's fleet only saved 1 to 2 percent at peak-carbon hours.\n"
               "Takeaways: rotation is memory bound, so PCIe sets the ceiling. And the cost of a FLOP depends on where and for whom it runs.")
for i, (h_, bl) in enumerate([
        ("Task 1", ["Limit: Colab has 2 vCPUs, no GDS test",
                    "Next: nvJPEG + GDS + streams",
                    "Takeaway: memory bound, PCIe sets the ceiling"]),
        ("Task 2", ["Limit: model energy figures are estimates",
                    "Ours: 0.71 J per 8K image on the T4",
                    "Takeaway: a FLOP's cost depends on where and for whom"])]):
    x = 0.6 + i * 6.2
    rect(sl, x, 1.5, 6.0, 5.3, fill=SKY)
    c = rect(sl, x, 1.5, 6.0, 0.85, fill=BLUE)
    label_in(c, h_, size=32, label="s14 head")
    text(sl, x + 0.15, 2.5, 5.7, 4.2, bl, size=30, bullets=True, label="s14 " + h_,
         gap=12)

# =====================================================================
# 15  Q&A
# =====================================================================
sl = prs.slides.add_slide(BLANK)
sl.background.fill.solid()
sl.background.fill.fore_color.rgb = NAVY
rect(sl, 0, 0, 0.4, 7.5, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
text(sl, 0.9, 2.2, 11.5, 1.6, "Questions?", size=80, bold=True, color=WHITE,
     label="qa", gap=0)
text(sl, 0.9, 4.2, 11.5, 1.2,
     ["Erwyna Soo Wen Xin and Taabish Farooq Bhat"], size=32, color=MID, label="qa names")
sl.notes_slide.notes_text_frame.text = (
    "ERWYNA (0:03)\nThank you. Any questions?\n\n"
    "Q&A RULE: verdict first, then the mechanism, then our number. Use parallel terms, no analogies. Name the problem before the fix.")

# =====================================================================
# APPENDIX
# =====================================================================
# A1
sl = new_slide("A1. Where the time goes (8K, 30 degrees)", "Taabish", appendix=True,
               sources="v2 nearest, pinned, Tesla T4, mean of 10 runs.",
               notes="APPENDIX (Taabish). Only for Q&A. H2D 8.06 ms, kernel 1.55 ms, D2H 7.58 ms, total 17.19 ms. CPU one core 509.3 ms. Mismatch 0 pixels.")
tot = 17.19
x0, wtot = 0.6, 12.1
segs = [("H2D", 8.06, BLUE), ("Kernel", 1.55, NAVY), ("D2H", 7.58, BLUE)]
xx = x0
for n_, v_, c_ in segs:
    w_ = wtot * v_ / tot
    r = rect(sl, xx, 1.9, w_ - 0.04, 1.2, fill=c_, radius=0.05)
    if v_ > 3:
        label_in(r, f"{n_} {v_} ms", size=30, label="a1 seg")
    xx += w_
text(sl, 0.6, 3.2, 12.1, 0.6, "Kernel is only 1.55 ms (9%)", size=30, bold=True,
     color=NAVY, align=PP_ALIGN.CENTER, label="a1 cap")
text(sl, 0.6, 4.1, 12.1, 2.8,
     ["Total on GPU: 17.19 ms", "CPU, one core: 509.3 ms", "Mismatch vs CPU: 0 pixels"],
     size=32, bullets=True, label="a1 bullets", gap=12)

# A2 pinned vs pageable + streams
sl = new_slide("A2. Pinned vs pageable, and streams", "Taabish", appendix=True,
               notes="APPENDIX (Taabish). Only for Q&A. Pageable 4.6 GB/s vs pinned 12.3 GB/s. Streams: 1.8x on a batch of 8 images, 1.36x on one image.")
picture(sl, os.path.join(T1, "results", "g3_pinned_vs_pageable.png"), 0.6, 1.8, 5.9, 3.3)
picture(sl, os.path.join(T1, "results", "g6_streams.png"), 6.8, 1.8, 5.9, 3.3)
text(sl, 0.6, 5.3, 12.1, 1.7,
     ["Pageable 4.6 GB/s vs pinned 12.3 GB/s",
      "Streams: 1.8x on 8 images, 1.36x on one image"],
     size=30, bullets=True, label="a2 bullets", gap=8)

# A2b code: streams
sl = new_slide("A2. Streams code (from our code)", "Taabish", appendix=True,
               sources="rotate_cuda.cu in our submission zip.",
               notes="APPENDIX (Taabish). Only for Q&A. Image i goes to stream i % 4. Async copies need pinned RAM. Same stream runs in order; different streams overlap.")
code_frame(sl, 0.6, 1.8, 12.1, 3.6, 3, "Round-robin streams loop with cudaMemcpyAsync",
           "rotate_cuda.cu lines 803 to 815")
text(sl, 0.6, 5.6, 12.1, 1.3,
     ["Image i goes to stream i % 4", "Async copies need pinned memory"],
     size=30, bullets=True, label="a2b bullets", gap=6)

# A3 sweeps
sl = new_slide("A3. Block sweep and image size sweep", "Taabish", appendix=True,
               notes="APPENDIX (Taabish). Only for Q&A. Nearest: flat from 64 to 256 threads per block (about 1.44 to 1.45 ms); 32x32 slower (1.88 ms) from register and cache pressure. Bilinear: 16x8 fastest (2.83 ms). End-to-end speed-up grows from 19x at 640x360 to 35x at 16K for nearest. Run to run the v2 kernel varied 1.45 to 2.11 ms; copies stayed 8.06 ms.")
picture(sl, os.path.join(T1, "results", "g4_block_sweep.png"), 0.6, 1.8, 6.0, 2.9)
picture(sl, os.path.join(T1, "results", "g5_size_sweep.png"), 6.8, 1.8, 6.0, 2.9)
text(sl, 0.6, 4.9, 12.1, 2.0,
     ["Nearest: block shape barely matters",
      "Bilinear: 16 x 8 is fastest (2.83 ms)",
      "Bigger images: bigger speed-up"],
     size=30, bullets=True, label="a3 bullets", gap=6)

# A4 timing code
sl = new_slide("A4. How we timed it (from our code)", "Taabish", appendix=True,
               sources="rotate_cuda.cu in our submission zip.",
               notes="APPENDIX (Taabish). Only for Q&A. Each stage is timed on the GPU clock with a pair of events. Warm-up run discarded, then mean of 10. Built with --fmad=false so v1 to v5 match the CPU byte for byte. v6 texture differs in 257,964 bytes (0.26%), each by at most 1, from the texture unit's low-precision filter weights.")
code_frame(sl, 0.6, 1.8, 6.4, 4.9, 4, "CUDA events around H2D, kernel, D2H",
           "rotate_cuda.cu lines 440 to 450")
text(sl, 7.2, 1.8, 5.6, 5.0,
     ["Events use the GPU clock", "Warm-up run discarded", "Mean of 10 runs",
      "--fmad=false: v1 to v5 match the CPU exactly"],
     size=30, bullets=True, label="a4 bullets", gap=12)

# A5 PCIe
sl = new_slide("A5. Why PCIe is the narrow pipe", "Taabish", appendix=True,
               sources="PCI-SIG (2019); NVIDIA (n.d.-d).",
               notes="APPENDIX (Taabish). Only for Q&A. T4 balance is about 25 FLOP per byte. Rotation does a few FLOPs per byte moved, so the kernel is memory bound. PCIe is about 20x narrower than GDDR6, so copies are 91 percent of the time for one image.")
bars = [("T4 GDDR6", 320, NAVY, "320 GB/s"), ("PCIe 3.0 x16", 15.75, BLUE, "15.75 GB/s"),
        ("Pinned copy", 12.3, BLUE, "12.3 GB/s")]
for i, (n_, v_, c_, l_) in enumerate(bars):
    y = 1.9 + i * 1.1
    text(sl, 0.6, y, 3.4, 0.9, n_, size=30, bold=True, color=NAVY,
         anchor=MSO_ANCHOR.MIDDLE, label="a5 name", gap=0)
    w_ = max(0.25, 5.6 * v_ / 320)
    rect(sl, 4.0, y + 0.1, w_, 0.7, fill=c_, radius=0.1)
    text(sl, 4.0 + w_ + 0.1, y, 3.4, 0.9, l_, size=30, anchor=MSO_ANCHOR.MIDDLE,
         label="a5 lbl", gap=0)
text(sl, 0.6, 5.3, 12.1, 1.6,
     ["PCIe is about 20x narrower than GDDR6",
      "T4 needs about 25 FLOP per byte; rotation does a few: memory bound"],
     size=30, bullets=True, label="a5 bullets", gap=8)

# A6 GDS
sl = new_slide("A6. GDS in practice", "Taabish", appendix=True,
               sources="NVIDIA (n.d.-a, n.d.-b, n.d.-c).",
               notes="APPENDIX (Taabish). Only for Q&A. cuFileDriverOpen, cuFileHandleRegister (file opened with O_DIRECT), cuFileBufRegister (GPU buffer), then cuFileRead into GPU memory. The CPU issues the calls; data DMAs from NVMe to GPU with no host bounce buffer. Without nvidia-fs and the rest, cuFile falls back to a host bounce buffer, which is what Colab would do.")
steps = ["cuFileDriverOpen", "cuFileHandleRegister", "cuFileBufRegister", "cuFileRead"]
for i, s_ in enumerate(steps):
    s = rect(sl, 0.6 + i * 3.1, 1.8, 2.9, 1.5, fill=BLUE if i < 3 else NAVY)
    label_in(s, s_, size=24, label="a6 step")
text(sl, 0.6, 3.6, 12.1, 3.3,
     ["Data goes NVMe to GPU, no host bounce buffer",
      "Needs nvidia-fs, O_DIRECT, aligned buffers",
      "Without them: falls back to a bounce buffer (Colab)"],
     size=30, bullets=True, label="a6 bullets", gap=10)

# A7 SIMD SIMT NPU
sl = new_slide("A7. SIMD vs SIMT vs NPU", "Taabish", appendix=True,
               sources="Lindholm et al. (2008); Jouppi et al. (2017).",
               notes="APPENDIX (Taabish). Only for Q&A. SIMD: one instruction on a wide register, lanes in lockstep, divergence is masked and wastes lanes. SIMT: warps of 32 threads run one instruction, each with its own index; divergent branches serialise inside the warp. NPU: fixed matrix-multiply units (systolic array) for low-precision inference, high TOPS per watt, little flexibility. Rotation is a gather, so it maps poorly.")
rows = [("SIMD (CPU)", "One instruction, wide register, lanes in lockstep"),
        ("SIMT (GPU)", "Warps of 32 threads, each with its own index"),
        ("NPU", "Fixed matrix units: fast, but rotation is a poor fit")]
for i, (h_, t_) in enumerate(rows):
    y = 1.9 + i * 1.65
    rect(sl, 0.6, y, 12.2, 1.5, fill=SKY)
    c = rect(sl, 0.6, y, 3.2, 1.5, fill=BLUE)
    label_in(c, h_, size=30, label="a7 chip")
    text(sl, 3.95, y, 8.7, 1.5, t_, size=30, anchor=MSO_ANCHOR.MIDDLE, label="a7 " + h_, gap=0)

# References (small text by necessity)
ref = open(os.path.join(ROOT, "references_APA7_slides.md")).read()
t1s = ref.index("## Task 1")
t2s = ref.index("## Task 2")


def refs(block):
    out = []
    for para in block.split("\n\n")[1:]:
        para = para.strip()
        if para and not para.startswith("#"):
            out.append(re.sub(r"\*", "", para).replace("–", "-").replace("—", "-"))
    return out


r1 = refs(ref[t1s:t2s])
r2 = refs(ref[t2s:])
half = (len(r2) + 1) // 2
for ttl, items in [("A8. References, Task 1 (APA 7)", r1),
                   ("A9. References, Task 2 (APA 7), part 1", r2[:half]),
                   ("A9. References, Task 2 (APA 7), part 2", r2[half:])]:
    sl = new_slide(ttl, "Taabish", appendix=True,
                   notes="APPENDIX. Full references. Small type is deliberate: this is a reference list, not a slide to read aloud.")
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(1.75), Inches(12.2), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        _run(p, it, 10, False, INK)
    tot = sum(n_lines(it, 10, 12.0) * 10 * 1.2 + 3 for it in items) / 72
    if tot > 5.5:
        WARN.append(f"{ttl}: refs need {tot:.1f}in")

out = os.path.join(HERE, "FIT3143_Applied2_Slides_v3.pptx")
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
for w in WARN:
    print("WARNING", w)
